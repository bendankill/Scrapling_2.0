"""V2.2.0 专项：中文产品字段、扩展信息映射与五级类目。"""

import copy
import csv
import json
import os

import pytest
from openpyxl import load_workbook

from crawler import EmagCrawler
from exporters import Exporters
from models import ProductItem
from output_schema import (
    CATEGORY_LEVEL_FIELDS,
    EXTRA_OUTPUT_FIELD_MAP,
    EXTRA_TOP_LEVEL_OUTPUT_FIELD_MAP,
    FAVORITE_DATA_OUTPUT_FIELD_MAP,
    PRODUCT_OUTPUT_FIELD_MAP,
    extract_category_levels,
    output_columns,
    product_to_output_dict,
    split_category_path,
    translate_extra,
)


FIXTURE_PATH = os.path.join(
    os.path.dirname(__file__), "fixtures", "v220_product_sample.json")
FORBIDDEN_FIELDS = {"好评率", "星级值", "规格详情", "详情描述"}


@pytest.fixture
def sample_dict():
    with open(FIXTURE_PATH, encoding="utf-8") as handle:
        return json.load(handle)


def _item(sample_dict, **overrides):
    data = copy.deepcopy(sample_dict)
    data.update(overrides)
    return ProductItem(**data)


class TestOutputFieldMapping:
    def test_complete_mapping_has_exact_34_fields(self):
        assert len(PRODUCT_OUTPUT_FIELD_MAP) == 34
        assert PRODUCT_OUTPUT_FIELD_MAP["pnk"] == "PNK码"
        assert PRODUCT_OUTPUT_FIELD_MAP["title"] == "产品标题"
        assert PRODUCT_OUTPUT_FIELD_MAP["price_current"] == "前端价格"
        assert PRODUCT_OUTPUT_FIELD_MAP["price_old"] == "PRP原价"
        assert PRODUCT_OUTPUT_FIELD_MAP["discount_percent"] == "前端折扣"
        assert PRODUCT_OUTPUT_FIELD_MAP["rating"] == "评论分数"
        assert PRODUCT_OUTPUT_FIELD_MAP["review_count"] == "评价数量"
        assert PRODUCT_OUTPUT_FIELD_MAP["main_image_url"] == "产品图片"
        assert PRODUCT_OUTPUT_FIELD_MAP["campaign_name"] == "链接打标"
        assert "活动名称" not in PRODUCT_OUTPUT_FIELD_MAP.values()

    def test_sample_values_and_types_are_unchanged(self, sample_dict):
        output = product_to_output_dict(sample_dict)
        for internal_name, output_name in PRODUCT_OUTPUT_FIELD_MAP.items():
            if internal_name == "extra":
                continue
            assert output[output_name] == sample_dict[internal_name]
            assert type(output[output_name]) is type(sample_dict[internal_name])
        assert output["品牌"] == ""
        assert output["活动价格"] is None
        assert set(PRODUCT_OUTPUT_FIELD_MAP).isdisjoint(output)

    def test_output_order_inserts_real_levels_after_source(self, sample_dict):
        output = product_to_output_dict(sample_dict)
        assert list(output) == output_columns(3)
        assert list(output)[3:6] == ["一级类", "二级类", "三级类"]

    def test_extra_mapping_preserves_unknown_values_and_structure(self, sample_dict):
        translated = product_to_output_dict(sample_dict)["扩展信息"]
        favorite = translated["收藏数据"]
        assert favorite["PNK码"] == "DR8D26BBM"
        assert favorite["产品ID"] == 12345678
        assert favorite["报价ID"] == 87654321
        assert favorite["是否存在系列商品"] is False
        assert favorite["是否为系列商品"] is False
        assert favorite["产品名称"] == sample_dict["extra"]["favorite_data"]["product_name"]
        assert favorite["货币"] == "RON"
        assert favorite["价格"] == 529.97
        assert favorite["类目路径"].endswith("/Boxe")
        assert favorite["options_modal"] == {"enabled": True}
        assert favorite["scm_super_category"] == {"id": 10, "name": "Audio"}
        assert translated["options_modal"] == {"mode": "technical"}
        assert translated["scm_super_category"] == {"id": 20, "name": "Electronics"}
        assert translated["data-has-unfair-price"] == "0"

    def test_extra_mapping_table_has_exact_confirmed_fields(self):
        assert len(EXTRA_OUTPUT_FIELD_MAP) == 16
        assert len(FAVORITE_DATA_OUTPUT_FIELD_MAP) == 9
        assert len(EXTRA_TOP_LEVEL_OUTPUT_FIELD_MAP) == 7
        assert EXTRA_OUTPUT_FIELD_MAP["favorite_data"] == "收藏数据"
        assert EXTRA_OUTPUT_FIELD_MAP["data-referrer"] == "来源路径"

    def test_favorite_only_key_at_extra_top_level_stays_original(self):
        output = product_to_output_dict({"extra": {"pnk": "TOP", "options_modal": {"id": 1}}})
        assert output["扩展信息"] == {"pnk": "TOP", "options_modal": {"id": 1}}

    def test_scenario_a_top_level_category_trails_do_not_collide(self):
        extra = {"category_trail": "A/B", "data-category-trail": "C/D"}
        translated = translate_extra(extra)
        assert translated == {"category_trail": "A/B", "类目路径": "C/D"}
        assert len(translated) == len(extra) == 2

    def test_scenario_b_existing_chinese_favorite_key_preserves_both_original_keys(self):
        extra = {"favorite_data": {"pnk": "RAW", "PNK码": "EXISTING"}}
        translated = translate_extra(extra)
        assert translated == {
            "收藏数据": {"pnk": "RAW", "PNK码": "EXISTING"},
        }
        assert len(translated["收藏数据"]) == 2

    def test_scenario_c_existing_chinese_top_key_preserves_both_original_keys(self):
        extra = {"data-category-name": "RawName", "类目名称": "ExistingName"}
        translated = translate_extra(extra)
        assert translated == {
            "data-category-name": "RawName",
            "类目名称": "ExistingName",
        }
        assert len(translated) == 2

    def test_scenario_d_same_label_in_different_layers_translates_normally(self):
        extra = {
            "favorite_data": {"category_trail": "A/B"},
            "data-category-trail": "C/D",
        }
        assert translate_extra(extra) == {
            "收藏数据": {"类目路径": "A/B"},
            "类目路径": "C/D",
        }

    def test_extra_translation_preserves_order_types_counts_and_input(self):
        extra = {
            "availability_id": 7,
            "unknown": [1, False, None, {"nested": "ă"}],
            "favorite_data": {
                "price": 529.97,
                "options_modal": {"enabled": True},
            },
        }
        before = copy.deepcopy(extra)
        translated = translate_extra(extra)
        assert list(translated) == ["库存状态ID", "unknown", "收藏数据"]
        assert list(translated["收藏数据"]) == ["价格", "options_modal"]
        assert len(translated) == len(extra)
        assert len(translated["收藏数据"]) == len(extra["favorite_data"])
        assert translated["unknown"] == extra["unknown"]
        assert type(translated["库存状态ID"]) is int
        assert type(translated["收藏数据"]["价格"]) is float
        assert extra == before

    def test_product_item_legacy_column_api_uses_unique_schema(self):
        assert ProductItem.csv_columns(3) == output_columns(3)
        assert ProductItem.field_names(5) == output_columns(5)
        assert [name for name, _ in ProductItem.excel_columns(2)] == output_columns(2)


class TestCategoryHierarchy:
    def test_three_levels_split_only_on_slash(self):
        path = "TV, Audio-Video & Foto/Audio HI-FI & Profesionale/Boxe"
        assert split_category_path(path) == [
            "TV, Audio-Video & Foto", "Audio HI-FI & Profesionale", "Boxe"]
        output = product_to_output_dict({
            "category_name": "Boxe", "extra": {"favorite_data": {"category_trail": path}}})
        assert output["一级类"] == "TV, Audio-Video & Foto"
        assert output["二级类"] == "Audio HI-FI & Profesionale"
        assert output["三级类"] == "Boxe"
        assert "四级类" not in output and "五级类" not in output
        assert "Audio Hi-Fi" not in output.values()

    @pytest.mark.parametrize("level_count", [1, 2, 3, 4, 5, 6])
    def test_one_to_six_level_paths(self, level_count):
        path = "/".join(f"L{index}" for index in range(1, level_count + 1))
        expected = [f"L{index}" for index in range(1, min(level_count, 5) + 1)]
        assert split_category_path(path) == expected

    def test_whitespace_empty_segments_comma_and_romanian_characters(self):
        path = "  Casă, Grădină // Îngrijire copii / Jucării  "
        assert split_category_path(path) == [
            "Casă, Grădină", "Îngrijire copii", "Jucării"]

    def test_missing_path_never_uses_category_name_as_guess(self):
        output = product_to_output_dict({"category_name": "Boxe", "extra": {}})
        assert not set(CATEGORY_LEVEL_FIELDS).intersection(output)

    def test_more_complete_consistent_source_wins_without_concatenation(self):
        extra = {
            "favorite_data": {"category_trail": "Electronice/Audio/Boxe"},
            "data-category-trail": "IT & Electronice/TV & Audio/Audio HI-FI/Boxe",
        }
        assert extract_category_levels(extra, "Boxe") == [
            "IT & Electronice", "TV & Audio", "Audio HI-FI", "Boxe"]

    def test_mismatched_favorite_source_falls_back_to_consistent_source(self):
        extra = {
            "favorite_data": {"category_trail": "Electronice/Televizoare"},
            "data-category-trail": "Electronice/Audio/Boxe",
        }
        assert extract_category_levels(extra, "Boxe") == ["Electronice", "Audio", "Boxe"]

    def test_equal_length_consistent_sources_use_favorite_priority(self):
        extra = {
            "favorite_data": {"category_trail": "Electronice/Audio/Boxe"},
            "data-category-trail": "Marketplace/Sunet/Boxe",
        }
        assert extract_category_levels(extra, "Boxe") == ["Electronice", "Audio", "Boxe"]

    def test_all_mismatched_sources_produce_no_guessed_levels(self):
        extra = {
            "favorite_data": {"category_trail": "Electronice/Televizoare"},
            "data-category-trail": "Marketplace/Casti",
        }
        assert extract_category_levels(extra, "Boxe") == []


class TestExportFormats:
    def test_json_chinese_keys_dynamic_levels_and_whitelist(self, tmp_path, sample_dict):
        exporters = Exporters(str(tmp_path))
        exporters.add_product(_item(sample_dict))
        exporters.finalize()
        with open(tmp_path / "products.json", encoding="utf-8") as handle:
            records = json.load(handle)
        assert len(records) == 1
        record = records[0]
        allowed = set(PRODUCT_OUTPUT_FIELD_MAP.values()) | set(CATEGORY_LEVEL_FIELDS)
        assert set(record) <= allowed
        assert set(record).isdisjoint(FORBIDDEN_FIELDS)
        assert set(PRODUCT_OUTPUT_FIELD_MAP).isdisjoint(record)
        assert "三级类" in record and "四级类" not in record and "五级类" not in record
        assert record["产品标题"] == sample_dict["title"]
        assert record["链接打标"] == sample_dict["campaign_name"]
        assert record["商品标签"] == sample_dict["badges"]
        assert "活动名称" not in record
        assert record["活动价格"] is None
        assert isinstance(record["前端价格"], float)
        assert isinstance(record["评价数量"], int)
        assert isinstance(record["扩展信息"], dict)

    def test_csv_utf8_bom_dynamic_columns_empty_cells_and_count(self, tmp_path, sample_dict):
        second = copy.deepcopy(sample_dict)
        second.update({"category_name": "Final", "pnk": "SECOND"})
        second["extra"]["favorite_data"]["category_trail"] = "L1/L2/L3/Final"
        exporters = Exporters(str(tmp_path))
        exporters.add_products([_item(sample_dict), ProductItem(**second)])
        exporters.finalize()

        raw = (tmp_path / "products.csv").read_bytes()
        assert raw.startswith(b"\xef\xbb\xbf")
        with open(tmp_path / "products.csv", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        assert len(rows) == 2
        assert rows[0].keys() == dict.fromkeys(output_columns(4)).keys()
        by_pnk = {row["PNK码"]: row for row in rows}
        assert by_pnk["DR8D26BBM"]["四级类"] == ""
        assert by_pnk["SECOND"]["四级类"] == "Final"
        assert by_pnk["DR8D26BBM"]["前端价格"] == "529.97"
        assert "五级类" not in rows[0]
        assert "Boxă" in by_pnk["DR8D26BBM"]["产品标题"]
        assert by_pnk["DR8D26BBM"]["链接打标"] == sample_dict["campaign_name"]
        assert "活动名称" not in rows[0]

    def test_xlsx_chinese_dynamic_columns_numeric_and_empty_values(self, tmp_path, sample_dict):
        second = copy.deepcopy(sample_dict)
        second.update({"category_name": "Final", "pnk": "SECOND", "price_promo": 499.5})
        second["extra"]["favorite_data"]["category_trail"] = "L1/L2/L3/L4/L5/Final"
        exporters = Exporters(str(tmp_path))
        exporters.add_products([_item(sample_dict), ProductItem(**second)])
        exporters.finalize()

        workbook = load_workbook(tmp_path / "products.xlsx", data_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        assert headers == output_columns(5)
        assert sheet.max_row - 1 == 2
        pnk_col = headers.index("PNK码") + 1
        rows = {sheet.cell(row, pnk_col).value: row for row in range(2, sheet.max_row + 1)}
        first_row = rows["DR8D26BBM"]
        assert sheet.cell(first_row, headers.index("五级类") + 1).value is None
        assert isinstance(sheet.cell(first_row, headers.index("前端价格") + 1).value, float)
        assert sheet.cell(first_row, headers.index("活动价格") + 1).value is None
        assert isinstance(sheet.cell(first_row, headers.index("评论分数") + 1).value, float)
        assert isinstance(sheet.cell(first_row, headers.index("评价数量") + 1).value, int)
        assert sheet.cell(first_row, headers.index("链接打标") + 1).value == sample_dict["campaign_name"]
        assert "活动名称" not in headers
        workbook.close()

    def test_three_formats_round_trip_campaign_and_collision_values(self, tmp_path, sample_dict):
        sample = copy.deepcopy(sample_dict)
        sample["campaign_name"] = "Top Favorite"
        sample["extra"] = {
            "category_trail": "A/B",
            "data-category-trail": "C/D",
            "favorite_data": {"pnk": "RAW", "PNK码": "EXISTING"},
            "data-category-name": "RawName",
            "类目名称": "ExistingName",
        }
        exporters = Exporters(str(tmp_path))
        exporters.add_product(ProductItem(**sample))
        exporters.finalize()

        with open(tmp_path / "products.json", encoding="utf-8") as handle:
            json_row = json.load(handle)[0]
        with open(tmp_path / "products.csv", encoding="utf-8-sig", newline="") as handle:
            csv_row = next(csv.DictReader(handle))
        workbook = load_workbook(tmp_path / "products.xlsx", data_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        xlsx_row = {header: sheet.cell(2, index + 1).value for index, header in enumerate(headers)}
        workbook.close()

        assert json_row["链接打标"] == csv_row["链接打标"] == xlsx_row["链接打标"] == "Top Favorite"
        assert "活动名称" not in json_row and "活动名称" not in csv_row and "活动名称" not in xlsx_row
        expected_extra = {
            "category_trail": "A/B",
            "类目路径": "C/D",
            "收藏数据": {"pnk": "RAW", "PNK码": "EXISTING"},
            "data-category-name": "RawName",
            "类目名称": "ExistingName",
        }
        assert json_row["扩展信息"] == expected_extra
        assert json.loads(csv_row["扩展信息"]) == expected_extra
        assert json.loads(xlsx_row["扩展信息"]) == expected_extra

    def test_three_format_counts_order_and_core_values_match(self, tmp_path, sample_dict):
        exporters = Exporters(str(tmp_path))
        products = [
            _item(sample_dict, category_name="Zed", pnk="Z", title="Titlu Z"),
            _item(sample_dict, category_name="Alpha", pnk="A", title="Titlu A"),
        ]
        for product in products:
            product.extra["favorite_data"]["category_trail"] = product.category_name
        exporters.add_products(products)
        internal_before = copy.deepcopy(exporters.get_products_sorted())
        exporters.finalize()

        with open(tmp_path / "products.json", encoding="utf-8") as handle:
            json_rows = json.load(handle)
        with open(tmp_path / "products.csv", encoding="utf-8-sig", newline="") as handle:
            csv_rows = list(csv.DictReader(handle))
        workbook = load_workbook(tmp_path / "products.xlsx", read_only=True, data_only=True)
        xlsx_count = workbook.active.max_row - 1
        workbook.close()

        assert len(json_rows) == len(csv_rows) == xlsx_count == 2
        assert [row["PNK码"] for row in json_rows] == ["A", "Z"]
        assert [row["产品标题"] for row in json_rows] == ["Titlu A", "Titlu Z"]
        assert exporters.get_products_sorted() == internal_before
        assert all("category_name" in item and "类目名称" not in item for item in internal_before)


class TestInternalBoundaryRegression:
    def test_product_item_and_exporter_internal_fields_stay_english(self, tmp_path, sample_dict):
        product = _item(sample_dict)
        internal = product.to_dict()
        assert internal["pnk"] == "DR8D26BBM"
        assert internal["main_image_local_path"] == "images/DR8D26BBM.jpg"
        assert "PNK码" not in internal and "本地图片路径" not in internal

        exporters = Exporters(str(tmp_path))
        exporters.add_product(product)
        assert exporters.get_products_sorted()[0]["main_image_local_path"] == "images/DR8D26BBM.jpg"
        assert exporters.get_json_buffer()[0]["本地图片路径"] == "images/DR8D26BBM.jpg"

    def test_run_summary_and_errors_schema_are_not_product_translated(self, tmp_path):
        output = tmp_path / "crawler"
        crawler = EmagCrawler(
            str(output), download_images=False, page_workers=1,
            category_workers=1, max_in_flight=2)
        summary = crawler.finalize()
        assert summary["version"] == "2.2.0"
        assert {"version", "status", "categories", "totals"} <= set(summary)
        assert "版本" not in summary and "状态" not in summary
        with open(output / "errors.csv", encoding="utf-8-sig", newline="") as handle:
            header = next(csv.reader(handle))
        assert header == crawler.ERROR_FIELDNAMES
