"""Follow-up S1 tests for origin validation, trust ordering, and HTTP sessions."""

import copy
import csv
import http.server
import json
import socketserver
import threading
from concurrent.futures import ThreadPoolExecutor

import pytest
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from category_hierarchy import (
    CategoryLevelRegistry,
    CategoryPathEvidence,
    extract_page_category_evidence,
)
from crawler import EmagCrawler
from exporters import Exporters
from models import ProductItem
from output_schema import (
    CATEGORY_LEVEL_FIELDS,
    PRODUCT_OUTPUT_FIELD_MAP,
    product_to_output_dict,
)


BOXE_URL = "https://www.emag.ro/boxe/c?ref=bc"
FULL_LEVELS = [
    "TV, Audio-Video & Foto",
    "Audio HI-FI & Profesionale",
    "Audio Hi-Fi",
    "Boxe",
]


def _soup(html):
    return BeautifulSoup(html, "lxml")


def _visible(levels, *, base="", leaf_href=None, attributes=None):
    items = ["<li>eMAG</li>"]
    attributes = attributes or {}
    for index, level in enumerate(levels):
        attrs = attributes.get(index, "")
        if index == len(levels) - 1:
            href = leaf_href
        else:
            href = f"{base}/level-{index}/c"
        if href is None:
            items.append(f"<li {attrs}>{level}</li>")
        else:
            items.append(f'<li {attrs}><a href="{href}">{level}</a></li>')
    return '<nav aria-label="Breadcrumb"><ol>' + "".join(items) + "</ol></nav>"


def _final(levels, source="visible_breadcrumb"):
    return CategoryPathEvidence(
        tuple(levels), source, 400,
        category_links_verified=True,
        validation_reason="verified category links",
    )


def _structured(levels):
    return CategoryPathEvidence(
        tuple(levels), "json_ld_breadcrumb", 350,
        structured_breadcrumb=True,
        validation_reason="verified structured breadcrumb",
    )


def _other_final(levels):
    return CategoryPathEvidence(
        tuple(levels), "product_path", 280,
        validation_reason="independent final evidence",
    )


def _tentative(levels):
    return CategoryPathEvidence(
        tuple(levels), "embedded_category_state", 150,
        is_tentative=True,
        validation_reason="unverified embedded path",
    )


class TestStrictBreadcrumbOriginAndIdentity:
    def test_plain_help_text_ending_in_current_category_is_rejected(self):
        html = """<nav aria-label="Breadcrumb"><ol><li>eMAG</li>
          <li>Contul meu</li><li>Ajutor</li><li>Boxe</li></ol></nav>"""
        assert extract_page_category_evidence(
            _soup(html), "Boxe", BOXE_URL) is None

    def test_account_help_links_with_data_id_are_rejected(self):
        html = """<nav aria-label="Breadcrumb"><ol><li>eMAG</li>
          <li data-id="10"><a href="/account">Contul meu</a></li>
          <li data-id="11"><a href="/help">Ajutor</a></li><li>Boxe</li>
        </ol></nav>"""
        assert extract_page_category_evidence(
            _soup(html), "Boxe", BOXE_URL) is None

    def test_generic_data_id_is_not_category_id(self):
        html = """<nav aria-label="Breadcrumb"><ol><li>eMAG</li>
          <li data-id="10">Audio</li><li data-id="11">Boxe</li>
        </ol></nav>"""
        assert extract_page_category_evidence(
            _soup(html), "Boxe", BOXE_URL) is None

    def test_explicit_category_id_without_link_is_allowed(self):
        html = """<nav aria-label="Breadcrumb"><ol><li>eMAG</li>
          <li data-category-id="10">Audio</li>
          <li category-id="11">Boxe</li></ol></nav>"""
        evidence = extract_page_category_evidence(
            _soup(html), "Boxe", BOXE_URL)
        assert list(evidence.levels) == ["Audio", "Boxe"]
        assert evidence.category_links_verified is True

    def test_invalid_link_cannot_be_hidden_by_category_id(self):
        html = """<nav aria-label="Breadcrumb"><ol><li>eMAG</li>
          <li data-category-id="10"><a href="/account">Audio</a></li>
          <li>Boxe</li></ol></nav>"""
        assert extract_page_category_evidence(
            _soup(html), "Boxe", BOXE_URL) is None

    def test_external_domain_category_shape_is_rejected(self):
        html = _visible(
            ["Audio", "Boxe"],
            base="https://evil.example",
            leaf_href="https://evil.example/boxe/c",
        )
        assert extract_page_category_evidence(
            _soup(html), "Boxe", BOXE_URL) is None

    def test_relative_category_urls_are_accepted(self):
        html = _visible(["Audio", "Boxe"], leaf_href="/boxe/c?ref=x")
        evidence = extract_page_category_evidence(
            _soup(html), "Boxe", BOXE_URL)
        assert list(evidence.levels) == ["Audio", "Boxe"]

    def test_same_origin_absolute_category_urls_are_accepted(self):
        html = _visible(
            ["Audio", "Boxe"],
            base="https://www.emag.ro",
            leaf_href="https://www.emag.ro/boxe/c?ref=x",
        )
        evidence = extract_page_category_evidence(
            _soup(html), "Boxe", BOXE_URL)
        assert list(evidence.levels) == ["Audio", "Boxe"]

    def test_emag_www_alias_is_accepted(self):
        html = _visible(
            ["Audio", "Boxe"],
            base="https://emag.ro",
            leaf_href="https://emag.ro/boxe/c",
        )
        evidence = extract_page_category_evidence(
            _soup(html), "Boxe", BOXE_URL)
        assert list(evidence.levels) == ["Audio", "Boxe"]

    def test_leaf_text_matches_but_leaf_url_other_category_is_rejected(self):
        html = _visible(["Audio", "Boxe"], leaf_href="/casti/c")
        assert extract_page_category_evidence(
            _soup(html), "Boxe", BOXE_URL) is None

    def test_pagination_and_query_are_ignored_for_leaf_identity(self):
        html = _visible(["Audio", "Boxe"], leaf_href="/boxe/c?ref=crumb")
        evidence = extract_page_category_evidence(
            _soup(html), "Boxe", "https://www.emag.ro/boxe/p2/c?ref=page")
        assert list(evidence.levels) == ["Audio", "Boxe"]

    def test_invalid_page_evidence_falls_back_to_favorite_data(self):
        invalid = extract_page_category_evidence(
            _soup(_visible(["Ajutor", "Boxe"], leaf_href=None)),
            "Boxe", BOXE_URL)
        output = product_to_output_dict({
            "category_name": "Boxe",
            "extra": {"favorite_data": {"category_trail": "TV/Audio/Boxe"}},
        }, category_evidence=invalid)
        assert [output[field] for field in CATEGORY_LEVEL_FIELDS[:3]] == [
            "TV", "Audio", "Boxe"]

    def test_missing_favorite_falls_back_to_data_category_trail(self):
        output = product_to_output_dict({
            "category_name": "Boxe",
            "extra": {"data-category-trail": "TV/Audio/Boxe"},
        })
        assert [output[field] for field in CATEGORY_LEVEL_FIELDS[:3]] == [
            "TV", "Audio", "Boxe"]

    def test_all_invalid_sources_emit_no_category_levels(self):
        invalid = extract_page_category_evidence(
            _soup("""<nav aria-label="Breadcrumb"><ol><li>eMAG</li>
              <li>Contul meu</li><li>Ajutor</li><li>Boxe</li></ol></nav>"""),
            "Boxe", BOXE_URL)
        output = product_to_output_dict({
            "category_name": "Boxe",
            "extra": {
                "favorite_data": {"category_trail": "TV/Televizoare"},
                "data-category-trail": "Marketplace/Casti",
            },
        }, category_evidence=invalid)
        assert not set(CATEGORY_LEVEL_FIELDS).intersection(output)

    def test_rejected_plain_text_never_enters_global_registry(self, tmp_path):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        evidence, _ = crawler._get_or_extract_category_evidence(
            _soup("""<nav aria-label="Breadcrumb"><ol><li>eMAG</li>
              <li>Contul meu</li><li>Ajutor</li><li>Boxe</li></ol></nav>"""),
            "Boxe", BOXE_URL)
        assert evidence is None
        assert crawler.exporters.get_category_level_evidence(BOXE_URL) is None


@pytest.mark.parametrize(
    "first,second,expected",
    [
        (_tentative(["Audio", "Boxe"]), _final(FULL_LEVELS), FULL_LEVELS),
        (_tentative(["TV", "Marketplace", "Audio", "Boxe"]),
         _final(["Audio", "Boxe"]), ["Audio", "Boxe"]),
        (_final(["Audio", "Boxe"]),
         _tentative(["TV", "Marketplace", "Audio", "Boxe"]),
         ["Audio", "Boxe"]),
        (_final(FULL_LEVELS), _tentative(["Audio", "Boxe"]), FULL_LEVELS),
        (_structured(["TV", "Audio", "Boxe"]),
         _structured(["TV", "Audio", "Audio Hi-Fi", "Boxe"]),
         ["TV", "Audio", "Audio Hi-Fi", "Boxe"]),
        (_other_final(["TV", "Audio", "Boxe"]),
         _other_final(["Marketplace", "Audio", "Boxe"]),
         ["TV", "Audio", "Boxe"]),
    ],
)
def test_registry_trust_replacement_matrix(first, second, expected):
    registry = CategoryLevelRegistry()
    registry.register(BOXE_URL, first)
    registry.register(BOXE_URL, second)
    assert list(registry.get(BOXE_URL).levels) == expected


class TestCacheTrustCrawlerAndConcurrency:
    def test_conflicting_verified_paths_do_not_replace_by_length(self):
        registry = CategoryLevelRegistry()
        first = _final(["TV", "Audio", "Boxe"])
        longer = _final(["Marketplace", "Recomandari", "Audio", "Boxe"])
        assert registry.register(BOXE_URL, first)
        assert not registry.register(BOXE_URL, longer)
        assert registry.get(BOXE_URL) == first

    def test_concurrent_random_order_converges_to_verified_evidence(self):
        registry = CategoryLevelRegistry()
        candidates = [
            _tentative(["TV", "Marketplace", "Audio", "Boxe"]),
            _final(["Audio", "Boxe"]),
            _tentative(["Audio", "Boxe"]),
        ] * 40
        with ThreadPoolExecutor(max_workers=12) as executor:
            list(executor.map(lambda item: registry.register(BOXE_URL, item), candidates))
        cached = registry.get(BOXE_URL)
        assert list(cached.levels) == ["Audio", "Boxe"]
        assert cached.is_tentative is False

    def test_crawler_tentative_long_upgrades_to_final_short(self, tmp_path):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        long_payload = json.dumps({
            "catalog": {"categoryPath": "TV/Marketplace/Audio/Boxe"}})
        first = _soup(f'<script id="__INITIAL_STATE__">{long_payload}</script>')
        second = _soup(_visible(["Audio", "Boxe"], leaf_href="/boxe/c"))
        crawler._get_or_extract_category_evidence(first, "Boxe", BOXE_URL)
        result, from_cache = crawler._get_or_extract_category_evidence(
            second, "Boxe", "https://www.emag.ro/boxe/p2/c")
        assert list(result.levels) == ["Audio", "Boxe"]
        assert result.is_tentative is False
        assert from_cache is False

    def test_final_cache_skips_later_tentative_page(self, tmp_path):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        crawler.exporters.register_category_levels(
            BOXE_URL, _final(["Audio", "Boxe"]))
        result, from_cache = crawler._get_or_extract_category_evidence(
            _soup("<html></html>"), "Boxe", "https://www.emag.ro/boxe/p2/c")
        assert list(result.levels) == ["Audio", "Boxe"]
        assert from_cache is True

    def test_third_upgrade_scan_can_install_final_evidence(self, tmp_path):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        payload = json.dumps({"catalog": {"categoryPath": "Audio/Boxe"}})
        tentative = _soup(f'<script id="__INITIAL_STATE__">{payload}</script>')
        crawler._get_or_extract_category_evidence(tentative, "Boxe", BOXE_URL)
        crawler._get_or_extract_category_evidence(
            tentative, "Boxe", "https://www.emag.ro/boxe/p2/c")
        final_page = _soup(_visible(FULL_LEVELS, leaf_href="/boxe/c"))
        result, _ = crawler._get_or_extract_category_evidence(
            final_page, "Boxe", "https://www.emag.ro/boxe/p3/c")
        assert list(result.levels) == FULL_LEVELS
        assert crawler._category_level_upgrade_checks[
            "emag.ro/boxe/c"] == 3

    def test_different_category_urls_remain_isolated(self):
        registry = CategoryLevelRegistry()
        registry.register(BOXE_URL, _final(["Audio", "Boxe"]))
        registry.register(
            "https://www.emag.ro/casti/c", _final(["Audio", "Casti"]))
        assert list(registry.get(BOXE_URL).levels) == ["Audio", "Boxe"]
        assert list(registry.get("https://www.emag.ro/casti/c").levels) == [
            "Audio", "Casti"]


class _SessionHandler(http.server.BaseHTTPRequestHandler):
    lock = threading.Lock()
    count = 0
    cookie_headers = []

    def do_GET(self):
        with self.lock:
            type(self).count += 1
            type(self).cookie_headers.append(self.headers.get("Cookie", ""))
        body = b"<html><body>ok</body></html>"
        self.send_response(200)
        if not self.headers.get("Cookie"):
            self.send_header("Set-Cookie", "codex_session=kept; Path=/")
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, _format, *_args):
        return


class _ThreadingServer(socketserver.ThreadingMixIn, socketserver.TCPServer):
    allow_reuse_address = True
    daemon_threads = True


@pytest.fixture
def session_server(monkeypatch):
    for name in ("ALL_PROXY", "HTTP_PROXY", "HTTPS_PROXY", "all_proxy",
                 "http_proxy", "https_proxy"):
        monkeypatch.delenv(name, raising=False)
    _SessionHandler.count = 0
    _SessionHandler.cookie_headers = []
    server = _ThreadingServer(("127.0.0.1", 0), _SessionHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_address[1]}/test/c"
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


class TestSessionRetryAndLifecycle:
    @pytest.mark.parametrize("invalid", [0, -1])
    def test_nonpositive_retries_are_rejected_clearly(self, tmp_path, invalid):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        crawler._session_config["retries"] = invalid
        with pytest.raises(ValueError, match=r"retries must be an integer >= 1"):
            crawler._get_client()
        assert crawler._all_sessions == []

    def test_retries_one_sends_exactly_one_local_http_request(
            self, tmp_path, session_server):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        crawler._session_config["retries"] = 1
        fetched = crawler._fetch_page(session_server)
        crawler._close_all_sessions()
        assert fetched.status == 200
        assert _SessionHandler.count == 1

    def test_same_thread_reuses_same_session(self, tmp_path, session_server):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        crawler._session_config["retries"] = 1
        first = crawler._get_client()
        assert first.get(session_server).status == 200
        second = crawler._get_client()
        assert second.get(session_server).status == 200
        crawler._close_all_sessions()
        assert first is second
        assert _SessionHandler.count == 2

    def test_different_threads_receive_isolated_sessions(
            self, tmp_path, session_server):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        crawler._session_config["retries"] = 1

        def request_once(_index):
            client = crawler._get_client()
            return id(client), client.get(session_server).status

        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(request_once, range(2)))
        crawler._close_all_sessions()
        assert len({client_id for client_id, _status in results}) == 2
        assert [status for _client_id, status in results] == [200, 200]

    def test_finalize_closes_all_session_managers(
            self, tmp_path, session_server):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        crawler._session_config["retries"] = 1
        assert crawler._get_client().get(session_server).status == 200
        managers = [manager for manager, _client in crawler._all_sessions]
        crawler.finalize()
        assert crawler._all_sessions == []
        assert all(not manager._is_alive for manager in managers)

    def test_cookie_continuity_is_preserved(self, tmp_path, session_server):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        crawler._session_config["retries"] = 1
        client = crawler._get_client()
        assert client.get(session_server).status == 200
        assert client.get(session_server).status == 200
        crawler._close_all_sessions()
        assert _SessionHandler.cookie_headers[0] == ""
        assert "codex_session=kept" in _SessionHandler.cookie_headers[1]


class TestThreeFormatTrustBoundary:
    def test_three_four_five_levels_round_trip_without_help_or_mutation(self, tmp_path):
        output_dir = tmp_path / "actual"
        output_dir.mkdir()
        exporter = Exporters(str(output_dir))
        specs = [
            ("Boxe", "boxe", ["TV", "Audio", "Boxe"]),
            ("Subwoofere", "subwoofere", ["TV", "Audio", "Hi-Fi", "Subwoofere"]),
            ("Accesorii", "accesorii", ["TV", "Audio", "Hi-Fi", "Boxe", "Accesorii"]),
        ]
        products = []
        before = []
        for index, (name, slug, levels) in enumerate(specs, 1):
            url = f"https://www.emag.ro/{slug}/c"
            product = ProductItem(
                category_name=name, category_url=url, source_page_url=url,
                page_number=1, position_in_page=index,
                product_id=str(index), pnk=f"PNK{index}", offer_id=f"OFF{index}",
                title=f"Produs {index}", product_url=f"https://www.emag.ro/p{index}/pd/PNK{index}/",
                price_current=100.25 + index, rating=4.5,
                review_count=index, campaign_name="Top Favorite",
                extra={"favorite_data": {"category_trail": "/".join(levels)}},
            )
            products.append(product)
            before.append(copy.deepcopy(product.to_dict()))
            exporter.register_category_levels(url, _final(levels))
            exporter.add_product(product)
        exporter.finalize()

        json_rows = json.loads((output_dir / "products.json").read_text(encoding="utf-8"))
        raw_csv = (output_dir / "products.csv").read_bytes()
        with (output_dir / "products.csv").open(
                encoding="utf-8-sig", newline="") as handle:
            csv_rows = list(csv.DictReader(handle))
        workbook = load_workbook(output_dir / "products.xlsx", data_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        xlsx_rows = [dict(zip(headers, [cell.value for cell in row]))
                     for row in sheet.iter_rows(min_row=2)]
        workbook.close()

        assert len(json_rows) == len(csv_rows) == len(xlsx_rows) == 3
        by_pnk = {row[PRODUCT_OUTPUT_FIELD_MAP["pnk"]]: row for row in json_rows}
        assert [field in by_pnk["PNK1"] for field in CATEGORY_LEVEL_FIELDS] == [
            True, True, True, False, False]
        assert [field in by_pnk["PNK2"] for field in CATEGORY_LEVEL_FIELDS] == [
            True, True, True, True, False]
        assert [field in by_pnk["PNK3"] for field in CATEGORY_LEVEL_FIELDS] == [
            True, True, True, True, True]
        assert headers[3:8] == list(CATEGORY_LEVEL_FIELDS)
        assert "Contul meu" not in json.dumps(json_rows, ensure_ascii=False)
        assert "Ajutor" not in json.dumps(json_rows, ensure_ascii=False)
        assert raw_csv.startswith(b"\xef\xbb\xbf")
        assert isinstance(json_rows[0][PRODUCT_OUTPUT_FIELD_MAP["extra"]], dict)
        assert isinstance(xlsx_rows[0][PRODUCT_OUTPUT_FIELD_MAP["price_current"]], float)
        assert isinstance(xlsx_rows[0][PRODUCT_OUTPUT_FIELD_MAP["rating"]], float)
        assert isinstance(xlsx_rows[0][PRODUCT_OUTPUT_FIELD_MAP["review_count"]], int)
        assert [product.to_dict() for product in products] == before
        allowed = set(PRODUCT_OUTPUT_FIELD_MAP.values()) | set(CATEGORY_LEVEL_FIELDS)
        assert all(set(row) <= allowed for row in json_rows)
