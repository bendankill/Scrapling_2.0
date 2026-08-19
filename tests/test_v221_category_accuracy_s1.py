"""S1 category-accuracy counterexamples for V2.2.1 production paths."""

import copy
import csv
import json
import threading
from concurrent.futures import ThreadPoolExecutor

from bs4 import BeautifulSoup
from openpyxl import load_workbook

import crawler as crawler_module
from category_hierarchy import (
    CategoryPathEvidence,
    category_paths_are_consistent,
    extract_page_category_evidence,
    select_best_category_evidence,
    select_product_category_levels,
)
from crawler import EmagCrawler
from exporters import Exporters
from models import ProductItem
from output_schema import CATEGORY_LEVEL_FIELDS, PRODUCT_OUTPUT_FIELD_MAP


FULL_LEVELS = [
    "TV, Audio-Video & Foto",
    "Audio HI-FI & Profesionale",
    "Audio Hi-Fi",
    "Boxe",
]


def _soup(html):
    return BeautifulSoup(html, "lxml")


def _category_breadcrumb(levels, *, include_leaf=True):
    selected = levels if include_leaf else levels[:-1]
    items = ["<li>eMAG</li>"]
    for index, level in enumerate(selected):
        if include_leaf and index == len(selected) - 1:
            items.append(f"<li>{level}</li>")
        else:
            items.append(f'<li><a href="/category-{index}/c">{level}</a></li>')
    return '<nav aria-label="Breadcrumb"><ol>' + "".join(items) + "</ol></nav>"


def _embedded_page(levels):
    payload = json.dumps({"catalog": {"categoryPath": "/".join(levels)}})
    return _soup(
        f'<html><body><script id="__INITIAL_STATE__">{payload}</script></body></html>')


def _final_evidence(levels=FULL_LEVELS, source="visible_breadcrumb"):
    return CategoryPathEvidence(
        tuple(levels), source, 400,
        category_links_verified=True,
        validation_reason="verified category links",
    )


def _tentative_evidence(levels, source="embedded_category_state"):
    return CategoryPathEvidence(
        tuple(levels), source, 150,
        is_tentative=True,
        validation_reason="unverified embedded path",
    )


class TestS1UnrelatedBreadcrumbs:
    def test_url_and_h1_do_not_authorize_help_breadcrumb_append(self):
        html = """<html><body><nav aria-label="Breadcrumb"><ol>
          <li>eMAG</li><li>Contul meu</li><li>Ajutor comenzi</li>
        </ol></nav><h1>Boxe</h1></body></html>"""
        evidence = extract_page_category_evidence(
            _soup(html), "Boxe", "https://www.emag.ro/boxe/c")
        assert evidence is None

    def test_real_category_breadcrumb_beats_longer_help_breadcrumb(self):
        correct = _category_breadcrumb(["TV", "Audio", "Boxe"])
        help_nav = """<nav aria-label="Breadcrumb"><ol><li>eMAG</li>
          <li>Contul meu</li><li>Ajutor</li><li>Comenzi</li>
          <li>Livrare</li></ol></nav>"""
        evidence = extract_page_category_evidence(
            _soup(f"<html><body>{correct}{help_nav}<h1>Boxe</h1></body></html>"),
            "Boxe", "https://www.emag.ro/boxe/c")
        assert list(evidence.levels) == ["TV", "Audio", "Boxe"]
        assert "Contul meu" not in evidence.levels

    def test_data_testid_text_alone_is_not_breadcrumb_semantics(self):
        html = """<html><body><div data-testid="account-breadcrumb">
          <span>Contul meu</span><span>Ajutor</span><span>Boxe</span>
        </div><h1>Boxe</h1></body></html>"""
        assert extract_page_category_evidence(
            _soup(html), "Boxe", "https://www.emag.ro/boxe/c") is None

    def test_footer_help_breadcrumb_is_rejected(self):
        html = """<html><body><footer><nav aria-label="Breadcrumb"><ol>
          <li>eMAG</li><li>Ajutor</li><li>Boxe</li>
        </ol></nav></footer><h1>Boxe</h1></body></html>"""
        assert extract_page_category_evidence(
            _soup(html), "Boxe", "https://www.emag.ro/boxe/c") is None

    def test_hidden_breadcrumb_is_rejected(self):
        html = """<html><body><div aria-hidden="TRUE"><nav aria-label="Breadcrumb">
          <ol><li>eMAG</li><li>TV</li><li>Boxe</li></ol>
        </nav></div><h1>Boxe</h1></body></html>"""
        assert extract_page_category_evidence(
            _soup(html), "Boxe", "https://www.emag.ro/boxe/c") is None

    def test_template_breadcrumb_is_rejected(self):
        html = """<html><body><template><nav aria-label="Breadcrumb"><ol>
          <li>eMAG</li><li>TV</li><li>Boxe</li>
        </ol></nav></template><h1>Boxe</h1></body></html>"""
        assert extract_page_category_evidence(
            _soup(html), "Boxe", "https://www.emag.ro/boxe/c") is None

    def test_plain_parent_text_cannot_append_current_category(self):
        html = """<html><body><nav aria-label="Breadcrumb"><ol>
          <li>eMAG</li><li>TV</li><li>Audio</li>
        </ol></nav><h1>Boxe</h1></body></html>"""
        assert extract_page_category_evidence(
            _soup(html), "Boxe", "https://www.emag.ro/boxe/c") is None

    def test_verified_category_parent_links_can_append_current_category(self):
        html = _category_breadcrumb(FULL_LEVELS, include_leaf=False)
        evidence = extract_page_category_evidence(
            _soup(f"<html><body>{html}<h1>Boxe</h1></body></html>"),
            "Boxe", "https://www.emag.ro/boxe/c")
        assert list(evidence.levels) == FULL_LEVELS
        assert evidence.current_category_appended is True
        assert evidence.category_links_verified is True

    def test_json_ld_category_urls_are_accepted(self):
        items = [
            {"@type": "ListItem", "position": index + 1, "name": name,
             "item": f"https://www.emag.ro/category-{index}/c"}
            for index, name in enumerate(["eMAG", *FULL_LEVELS])
        ]
        items[0]["item"] = "https://www.emag.ro/"
        items[-1]["item"] = "https://www.emag.ro/boxe/c"
        payload = {"@type": "BreadcrumbList", "itemListElement": items}
        html = f"<script type='application/ld+json'>{json.dumps(payload)}</script>"
        evidence = extract_page_category_evidence(
            _soup(html), "Boxe", "https://www.emag.ro/boxe/c")
        assert list(evidence.levels) == FULL_LEVELS
        assert evidence.structured_breadcrumb is True
        assert evidence.category_links_verified is True

    def test_json_ld_help_urls_are_rejected(self):
        payload = {
            "@type": "BreadcrumbList",
            "itemListElement": [
                {"position": 1, "name": "eMAG", "item": "https://www.emag.ro/"},
                {"position": 2, "name": "Contul meu", "item": "https://www.emag.ro/account"},
                {"position": 3, "name": "Boxe", "item": "https://www.emag.ro/help/boxe"},
            ],
        }
        html = f"<script type='application/ld+json'>{json.dumps(payload)}</script>"
        assert extract_page_category_evidence(
            _soup(html), "Boxe", "https://www.emag.ro/boxe/c") is None


class TestS1EvidenceConflictRules:
    def test_short_high_reliability_beats_long_conflicting_tentative_path(self):
        strong = _final_evidence(["TV", "Boxe"])
        weak = _tentative_evidence(["Marketplace", "Recomandari", "Boxe"])
        assert select_best_category_evidence([strong, weak]) == strong

    def test_long_consistent_tentative_extension_cannot_replace_final_path(self):
        strong = _final_evidence(["TV", "Audio", "Boxe"])
        extension = _tentative_evidence(["TV", "Audio", "Audio Hi-Fi", "Boxe"])
        assert select_best_category_evidence([strong, extension]) == strong

    def test_two_equally_strong_conflicting_paths_are_ambiguous(self):
        first = _final_evidence(["TV", "Audio", "Boxe"])
        second = _final_evidence(["Marketplace", "Audio", "Boxe"], "json_ld_breadcrumb")
        assert select_best_category_evidence([first, second]) is None

    def test_visible_and_jsonld_verified_conflict_is_ambiguous(self):
        visible = _category_breadcrumb(["TV", "Audio", "Boxe"])
        payload = {
            "@type": "BreadcrumbList",
            "itemListElement": [
                {"position": 1, "name": "eMAG", "item": "https://www.emag.ro/"},
                {"position": 2, "name": "Marketplace",
                 "item": "https://www.emag.ro/marketplace/c"},
                {"position": 3, "name": "Boxe",
                 "item": "https://www.emag.ro/boxe/c"},
            ],
        }
        script = f"<script type='application/ld+json'>{json.dumps(payload)}</script>"
        html = f"<html><head>{script}</head><body>{visible}</body></html>"
        assert extract_page_category_evidence(
            _soup(html), "Boxe", "https://www.emag.ro/boxe/c") is None

    def test_conflicting_embedded_paths_are_rejected(self):
        payload = {"one": {"categoryPath": "TV/Audio/Boxe"},
                   "two": {"categoryPath": "Marketplace/Recommended/Boxe"}}
        html = f'<script id="__INITIAL_STATE__">{json.dumps(payload)}</script>'
        assert extract_page_category_evidence(
            _soup(html), "Boxe", "https://www.emag.ro/boxe/c") is None

    def test_current_category_in_middle_is_rejected(self):
        evidence = _final_evidence(["TV", "Boxe", "Accesorii"])
        levels = select_product_category_levels(
            {}, "Boxe", category_evidence=evidence)
        assert levels == []

    def test_consistency_supports_ordered_missing_middle_levels(self):
        assert category_paths_are_consistent(
            ["TV", "Audio", "Boxe"],
            ["TV", "Audio", "Audio Hi-Fi", "Boxe"],
        )
        assert not category_paths_are_consistent(
            ["TV", "Video", "Boxe"],
            ["TV", "Audio", "Audio Hi-Fi", "Boxe"],
        )


class TestS1PageProductComparison:
    def test_two_level_page_does_not_overwrite_four_level_favorite_path(self):
        extra = {"favorite_data": {"category_trail": "/".join(FULL_LEVELS)}}
        levels = select_product_category_levels(
            extra, "Boxe", category_evidence=_final_evidence(["Audio", "Boxe"]))
        assert levels == FULL_LEVELS

    def test_four_level_page_completes_consistent_three_level_favorite_path(self):
        extra = {"favorite_data": {"category_trail":
                                    "TV, Audio-Video & Foto/Audio HI-FI & Profesionale/Boxe"}}
        levels = select_product_category_levels(
            extra, "Boxe", category_evidence=_final_evidence())
        assert levels == FULL_LEVELS

    def test_four_level_page_completes_consistent_card_path(self):
        extra = {"data-category-trail": "TV, Audio-Video & Foto/Boxe"}
        levels = select_product_category_levels(
            extra, "Boxe", category_evidence=_final_evidence())
        assert levels == FULL_LEVELS

    def test_tentative_conflicting_page_falls_back_to_favorite(self):
        extra = {"favorite_data": {"category_trail": "TV/Audio/Boxe"}}
        levels = select_product_category_levels(
            extra, "Boxe",
            category_evidence=_tentative_evidence(["Marketplace", "Recommended", "Boxe"]),
        )
        assert levels == ["TV", "Audio", "Boxe"]

    def test_invalid_page_falls_back_to_card_after_invalid_favorite(self):
        extra = {
            "favorite_data": {"category_trail": "TV/Televizoare"},
            "data-category-trail": "TV/Audio/Boxe",
        }
        invalid = _final_evidence(["TV", "Boxe", "Accesorii"])
        assert select_product_category_levels(
            extra, "Boxe", category_evidence=invalid) == ["TV", "Audio", "Boxe"]


class TestS1CrawlerCacheUpgrade:
    def test_real_crawler_chain_upgrades_tentative_two_to_final_four(self, tmp_path):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        first, first_cached = crawler._get_or_extract_category_evidence(
            _embedded_page(["Audio", "Boxe"]), "Boxe", "https://www.emag.ro/boxe/c")
        full_html = f"<html><body>{_category_breadcrumb(FULL_LEVELS)}</body></html>"
        second, second_cached = crawler._get_or_extract_category_evidence(
            _soup(full_html), "Boxe", "https://www.emag.ro/boxe/p2/c")
        assert list(first.levels) == ["Audio", "Boxe"]
        assert first_cached is False
        assert list(second.levels) == FULL_LEVELS
        assert second.is_tentative is False
        assert second_cached is False

    def test_final_four_level_cache_cannot_be_downgraded(self, tmp_path):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        full_html = f"<html><body>{_category_breadcrumb(FULL_LEVELS)}</body></html>"
        crawler._get_or_extract_category_evidence(
            _soup(full_html), "Boxe", "https://www.emag.ro/boxe/c")
        result, from_cache = crawler._get_or_extract_category_evidence(
            _embedded_page(["Audio", "Boxe"]), "Boxe", "https://www.emag.ro/boxe/p2/c")
        assert list(result.levels) == FULL_LEVELS
        assert from_cache is False

    def test_low_embedded_cache_upgrades_to_verified_breadcrumb(self, tmp_path):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        crawler._get_or_extract_category_evidence(
            _embedded_page(["Marketplace", "Boxe"]), "Boxe",
            "https://www.emag.ro/boxe/c")
        verified = _soup(
            f"<html><body>{_category_breadcrumb(FULL_LEVELS)}</body></html>")
        result, _ = crawler._get_or_extract_category_evidence(
            verified, "Boxe", "https://www.emag.ro/boxe/p2/c")
        assert list(result.levels) == FULL_LEVELS
        assert result.category_links_verified is True

    def test_final_cache_detects_later_final_conflict(self, tmp_path, monkeypatch):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        crawler.exporters.register_category_levels(
            "https://www.emag.ro/boxe/c", _final_evidence())
        calls = {"count": 0}

        def conflicting_decision(*_args):
            calls["count"] += 1
            from category_hierarchy import CategoryEvidenceDecision, CategoryEvidenceStatus
            return CategoryEvidenceDecision(
                CategoryEvidenceStatus.FINAL,
                _final_evidence(["Marketplace", "Boxe"]))

        monkeypatch.setattr(crawler_module, "extract_page_category_decision", conflicting_decision)
        result, from_cache = crawler._get_or_extract_category_evidence(
            _soup("<html></html>"), "Boxe", "https://www.emag.ro/boxe/p2/c",
            page_number=2)
        assert result is None
        assert from_cache is False
        assert calls["count"] == 1

    def test_concurrent_same_category_converges_on_best_evidence(self, tmp_path):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        tentative = _embedded_page(["Audio", "Boxe"])
        final = _soup(f"<html><body>{_category_breadcrumb(FULL_LEVELS)}</body></html>")
        jobs = [(tentative, 1), (final, 2)] * 20
        with ThreadPoolExecutor(max_workers=8) as executor:
            list(executor.map(
                lambda job: crawler._get_or_extract_category_evidence(
                    job[0], "Boxe", f"https://www.emag.ro/boxe/p{job[1]}/c"),
                jobs,
            ))
        cached = crawler.exporters.get_category_level_evidence(
            "https://www.emag.ro/boxe/c")
        assert list(cached.levels) == FULL_LEVELS
        assert cached.is_tentative is False

    def test_concurrent_categories_remain_isolated(self, tmp_path):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        boxe = _soup(f"<html><body>{_category_breadcrumb(FULL_LEVELS)}</body></html>")
        casti_levels = ["Electronice", "Audio", "Casti"]
        casti = _soup(f"<html><body>{_category_breadcrumb(casti_levels)}</body></html>")
        jobs = [(boxe, "Boxe", "boxe"), (casti, "Casti", "casti")] * 20
        with ThreadPoolExecutor(max_workers=8) as executor:
            list(executor.map(
                lambda job: crawler._get_or_extract_category_evidence(
                    job[0], job[1], f"https://www.emag.ro/{job[2]}/c"),
                jobs,
            ))
        assert list(crawler.exporters.get_category_level_evidence(
            "https://www.emag.ro/boxe/c").levels) == FULL_LEVELS
        assert list(crawler.exporters.get_category_level_evidence(
            "https://www.emag.ro/casti/c").levels) == casti_levels

    def test_tentative_cache_upgrade_scans_are_bounded(self, tmp_path, monkeypatch):
        crawler = EmagCrawler(str(tmp_path / "out"), download_images=False)
        calls = {"count": 0}
        original = crawler_module.extract_page_category_decision

        def counted(*args):
            calls["count"] += 1
            return original(*args)

        monkeypatch.setattr(crawler_module, "extract_page_category_decision", counted)
        soup = _embedded_page(["Audio", "Boxe"])
        for page in range(1, 20):
            crawler._get_or_extract_category_evidence(
                soup, "Boxe", f"https://www.emag.ro/boxe/p{page}/c",
                page_number=page)
        assert calls["count"] == crawler._category_level_max_upgrade_checks == 3


class TestS1ProductionWrite:
    def test_cache_upgrade_is_used_by_all_three_formats_without_mutation(self, tmp_path):
        sample = ProductItem(
            category_name="Boxe",
            category_url="https://www.emag.ro/boxe/c",
            source_page_url="https://www.emag.ro/boxe/c",
            page_number=1,
            position_in_page=1,
            product_id="123",
            pnk="DR8D26BBM",
            offer_id="456",
            title="Boxă portabilă Test",
            product_url="https://www.emag.ro/test/pd/DR8D26BBM/",
            price_current=529.97,
            rating=4.8,
            review_count=4,
            campaign_name="Top Favorite",
            extra={
                "favorite_data": {"category_trail": "/".join(FULL_LEVELS)},
                "data-category-trail": "/".join(FULL_LEVELS),
            },
        )
        before = copy.deepcopy(sample.to_dict())
        output_dir = tmp_path / "actual"
        output_dir.mkdir()
        exporter = Exporters(str(output_dir))
        exporter.register_category_levels(sample.category_url, _tentative_evidence(["Audio", "Boxe"]))
        exporter.register_category_levels(sample.category_url, _final_evidence())
        exporter.add_product(sample)
        exporter.finalize()

        json_rows = json.loads((output_dir / "products.json").read_text(encoding="utf-8"))
        raw_csv = (output_dir / "products.csv").read_bytes()
        with (output_dir / "products.csv").open(encoding="utf-8-sig", newline="") as handle:
            csv_rows = list(csv.DictReader(handle))
        workbook = load_workbook(output_dir / "products.xlsx", data_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        xlsx_rows = [dict(zip(headers, [cell.value for cell in row]))
                     for row in sheet.iter_rows(min_row=2)]
        workbook.close()

        assert len(json_rows) == len(csv_rows) == len(xlsx_rows) == 1
        assert [json_rows[0][field] for field in CATEGORY_LEVEL_FIELDS[:4]] == FULL_LEVELS
        assert CATEGORY_LEVEL_FIELDS[4] not in json_rows[0]
        assert "Contul meu" not in json.dumps(json_rows, ensure_ascii=False)
        assert headers[3:7] == list(CATEGORY_LEVEL_FIELDS[:4])
        assert raw_csv.startswith(b"\xef\xbb\xbf")
        assert isinstance(json_rows[0][PRODUCT_OUTPUT_FIELD_MAP["extra"]], dict)
        assert isinstance(xlsx_rows[0][PRODUCT_OUTPUT_FIELD_MAP["price_current"]], float)
        assert isinstance(xlsx_rows[0][PRODUCT_OUTPUT_FIELD_MAP["rating"]], float)
        assert isinstance(xlsx_rows[0][PRODUCT_OUTPUT_FIELD_MAP["review_count"]], int)
        assert sample.to_dict() == before
        allowed = set(PRODUCT_OUTPUT_FIELD_MAP.values()) | set(CATEGORY_LEVEL_FIELDS)
        assert set(json_rows[0]) <= allowed
