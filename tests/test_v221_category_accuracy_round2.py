"""V2.2.1 round-two accuracy, deterministic evidence, and session hardening."""

import copy
import csv
import hashlib
import html
import http.server
import json
import os
import random
import socketserver
import threading
import time
from concurrent.futures import ThreadPoolExecutor

import pytest
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from category_hierarchy import (
    CategoryEvidenceStatus,
    CategoryLevelRegistry,
    CategoryPathEvidence,
    _visible_breadcrumb_paths,
    decide_category_evidence,
    extract_page_category_decision,
    extract_page_category_evidence,
    select_product_category_levels,
)
from crawler import EmagCrawler, FetchResult
from exporters import Exporters
from models import ProductItem
from output_schema import CATEGORY_LEVEL_FIELDS, product_to_output_dict


BOXE_URL = "https://www.emag.ro/boxe/c?ref=bc"
FINAL_A = ["TV", "Audio", "Boxe"]
FINAL_B = ["Marketplace", "Recomandari", "Audio", "Boxe"]


def _soup(markup):
    return BeautifulSoup(markup, "lxml")


def _final(levels, source="visible_breadcrumb"):
    return CategoryPathEvidence(
        tuple(levels), source, 400,
        category_links_verified=True,
        validation_reason="verified category evidence",
    )


def _temporary(levels):
    return CategoryPathEvidence(
        tuple(levels), "embedded_category_state", 150,
        is_tentative=True,
        validation_reason="unverified embedded evidence",
    )


def _linked_breadcrumb(levels, *, current="Boxe"):
    items = ["<li>eMAG</li>"]
    for index, level in enumerate(levels):
        slug = "boxe" if level == current else f"level-{index}"
        items.append(f'<li><a href="/{slug}/c">{level}</a></li>')
    return '<nav aria-label="Breadcrumb"><ol>' + "".join(items) + "</ol></nav>"


SCHEMA_CASES = {
    "parent_item_link": (
        """<nav aria-label="Breadcrumb"><ol itemscope itemtype="https://schema.org/BreadcrumbList">
        <li itemprop="itemListElement"><a itemprop="item" href="/tv/c"><span itemprop="name">TV</span></a><meta itemprop="position" content="1"></li>
        <li itemprop="itemListElement"><a itemprop="item" href="/boxe/c"><span itemprop="name">Boxe</span></a><meta itemprop="position" content="2"></li></ol></nav>""",
        ["TV", "Boxe"],
    ),
    "item_is_link": (
        """<nav aria-label="Breadcrumb"><a href="/tv/c">TV</a><a href="/boxe/c">Boxe</a></nav>""",
        ["TV", "Boxe"],
    ),
    "name_contains_link": (
        """<nav aria-label="Breadcrumb"><ol><li itemprop="itemListElement"><span itemprop="name"><a href="/tv/c">TV</a></span></li>
        <li itemprop="itemListElement"><span itemprop="name"><a href="/boxe/c">Boxe</a></span></li></ol></nav>""",
        ["TV", "Boxe"],
    ),
    "name_and_item_are_siblings": (
        """<nav aria-label="Breadcrumb"><ol><li itemprop="itemListElement"><span itemprop="name">TV</span><a itemprop="item" href="/tv/c"></a></li>
        <li itemprop="itemListElement"><span itemprop="name">Boxe</span><a itemprop="item" href="/boxe/c"></a></li></ol></nav>""",
        ["TV", "Boxe"],
    ),
    "no_cross_item_parent_lookup": (
        """<nav aria-label="Breadcrumb"><a href="/help"><ol><li itemprop="itemListElement"><span itemprop="name">TV</span></li>
        <li itemprop="itemListElement"><span itemprop="name">Boxe</span></li></ol></a></nav>""",
        None,
    ),
    "outer_help_not_bound": (
        """<a href="/help">Ajutor</a><nav aria-label="Breadcrumb"><ol><li><a href="/tv/c">TV</a></li><li><a href="/boxe/c">Boxe</a></li></ol></nav>""",
        ["TV", "Boxe"],
    ),
    "outer_category_not_bound": (
        """<a href="/unrelated/c">Other</a><nav aria-label="Breadcrumb"><ol><li>TV</li><li>Boxe</li></ol></nav>""",
        None,
    ),
    "wrong_leaf_url": (
        """<nav aria-label="Breadcrumb"><ol><li><a href="/tv/c">TV</a></li><li><a href="/casti/c">Boxe</a></li></ol></nav>""",
        None,
    ),
    "json_ld_still_works": (
        '<script type="application/ld+json">' + json.dumps({
            "@type": "BreadcrumbList",
            "itemListElement": [
                {"position": 1, "name": "TV", "item": "/tv/c"},
                {"position": 2, "name": "Boxe", "item": "/boxe/c"},
            ],
        }) + "</script>",
        ["TV", "Boxe"],
    ),
    "legacy_visible_still_works": (_linked_breadcrumb(["TV", "Boxe"]), ["TV", "Boxe"]),
}


@pytest.mark.parametrize("case", list(SCHEMA_CASES))
def test_a_schema_and_visible_item_binding(case):
    """A1-A10: link lookup stays inside the matching logical breadcrumb item."""
    markup, expected = SCHEMA_CASES[case]
    evidence = extract_page_category_evidence(_soup(markup), "Boxe", BOXE_URL)
    assert (list(evidence.levels) if evidence else None) == expected
    if case == "parent_item_link":
        assert evidence.structured_breadcrumb is True
        assert evidence.category_links_verified is True


NODE_CASES = (
    "middle_duplicate_url", "middle_duplicate_id", "remove_home",
    "remove_empty", "position_sort", "validate_before_five_level_cut",
    "generic_data_id", "bad_link_masks_category_id",
)


@pytest.mark.parametrize("case", NODE_CASES)
def test_b_node_metadata_remains_bound(case):
    """B11-B18: names, URLs and IDs are cleaned as indivisible nodes."""
    if case in {"middle_duplicate_url", "middle_duplicate_id"}:
        markup = """<nav aria-label="Breadcrumb"><ol>
        <li><a href="/home">Home</a></li>
        <li data-category-id="tv"><a href="/tv/c">TV</a></li>
        <li data-category-id="audio-first"><a href="/audio/c">Audio</a></li>
        <li data-category-id="audio-second"><a href="/wrong-duplicate/c">Audio</a></li>
        <li data-category-id="boxe"><a href="/boxe/c">Boxe</a></li></ol></nav>"""
        descriptor = _visible_breadcrumb_paths(_soup(markup), BOXE_URL)[0]
        assert [node.name for node in descriptor.nodes] == ["TV", "Audio", "Boxe"]
        assert [node.url for node in descriptor.nodes] == ["/tv/c", "/audio/c", "/boxe/c"]
        assert [node.category_id for node in descriptor.nodes] == [
            "tv", "audio-first", "boxe"]
    elif case == "remove_home":
        descriptor = _visible_breadcrumb_paths(_soup(
            '<nav aria-label="Breadcrumb"><ol><li><a href="/">Acasă</a></li>'
            '<li><a href="/tv/c">TV</a></li><li><a href="/boxe/c">Boxe</a></li>'
            '</ol></nav>'), BOXE_URL)[0]
        assert [node.name for node in descriptor.nodes] == ["TV", "Boxe"]
        assert descriptor.nodes[0].url == "/tv/c"
    elif case == "remove_empty":
        descriptor = _visible_breadcrumb_paths(_soup(
            '<nav aria-label="Breadcrumb"><ol><li><a href="/empty/c"></a></li>'
            '<li><a href="/tv/c">TV</a></li><li><a href="/boxe/c">Boxe</a></li>'
            '</ol></nav>'), BOXE_URL)[0]
        assert [(node.name, node.url) for node in descriptor.nodes] == [
            ("TV", "/tv/c"), ("Boxe", "/boxe/c")]
    elif case == "position_sort":
        markup = """<nav aria-label="Breadcrumb"><ol>
        <li itemprop="itemListElement"><a itemprop="item" href="/boxe/c"><span itemprop="name">Boxe</span></a><meta itemprop="position" content="3"></li>
        <li itemprop="itemListElement"><a itemprop="item" href="/tv/c"><span itemprop="name">TV</span></a><meta itemprop="position" content="1"></li>
        <li itemprop="itemListElement"><a itemprop="item" href="/audio/c"><span itemprop="name">Audio</span></a><meta itemprop="position" content="2"></li></ol></nav>"""
        descriptor = _visible_breadcrumb_paths(_soup(markup), BOXE_URL)[0]
        assert [(node.name, node.url) for node in descriptor.nodes] == [
            ("TV", "/tv/c"), ("Audio", "/audio/c"), ("Boxe", "/boxe/c")]
    elif case == "validate_before_five_level_cut":
        levels = ["L1", "L2", "L3", "L4", "L5", "Boxe"]
        evidence = extract_page_category_evidence(
            _soup(_linked_breadcrumb(levels)), "Boxe", BOXE_URL)
        assert list(evidence.levels) == levels
        assert select_product_category_levels(
            {}, "Boxe", category_evidence=evidence) == levels[:5]
    elif case == "generic_data_id":
        markup = '<nav aria-label="Breadcrumb"><ol><li data-id="1">TV</li><li data-id="2">Boxe</li></ol></nav>'
        assert extract_page_category_evidence(_soup(markup), "Boxe", BOXE_URL) is None
    else:
        markup = '<nav aria-label="Breadcrumb"><ol><li data-category-id="1"><a href="/help">TV</a></li><li><a href="/boxe/c">Boxe</a></li></ol></nav>'
        assert extract_page_category_evidence(_soup(markup), "Boxe", BOXE_URL) is None


ORDER_CASES = (
    "category_path_first", "breadcrumb_first", "random_order", "identical_temporary",
    "consistent_temporary_extension", "temporary_conflict", "fallback_favorite",
    "fallback_card", "all_invalid", "final_beats_temporary", "final_conflict",
)


def _embedded_decision(payload):
    markup = '<script id="__INITIAL_STATE__" type="application/json">' + json.dumps(payload) + "</script>"
    return extract_page_category_decision(_soup(markup), "Boxe", BOXE_URL)


@pytest.mark.parametrize("case", ORDER_CASES)
def test_c_same_page_decision_is_order_independent(case):
    """C19-C29: collect first, decide once, and fall back deterministically."""
    conflict_a = ["TV", "Audio", "Boxe"]
    conflict_b = ["Marketplace", "Recommended", "Boxe"]
    if case in {"category_path_first", "breadcrumb_first"}:
        pairs = (["categoryPath", conflict_a], ["breadcrumb", conflict_b])
        if case == "breadcrumb_first":
            pairs = tuple(reversed(pairs))
        decision = _embedded_decision(dict(pairs))
        assert decision.status == CategoryEvidenceStatus.TEMPORARY_CONFLICTED
        assert decision.evidence is None
    elif case == "random_order":
        candidates = [_temporary(conflict_a), _temporary(conflict_b)]
        for seed in range(20):
            random.Random(seed).shuffle(candidates)
            assert decide_category_evidence(candidates).status == (
                CategoryEvidenceStatus.TEMPORARY_CONFLICTED)
    elif case == "identical_temporary":
        decision = decide_category_evidence([
            _temporary(conflict_a), _temporary(conflict_a)])
        assert decision.status == CategoryEvidenceStatus.TEMPORARY
        assert list(decision.evidence.levels) == conflict_a
    elif case == "consistent_temporary_extension":
        decision = decide_category_evidence([
            _temporary(["TV", "Boxe"]), _temporary(conflict_a)])
        assert list(decision.evidence.levels) == conflict_a
    elif case == "temporary_conflict":
        assert decide_category_evidence([
            _temporary(conflict_a), _temporary(conflict_b)]).evidence is None
    elif case == "fallback_favorite":
        extra = {"favorite_data": {"category_trail": "TV/Audio/Boxe"}}
        assert select_product_category_levels(extra, "Boxe") == conflict_a
    elif case == "fallback_card":
        extra = {"data-category-trail": "TV/Audio/Boxe"}
        assert select_product_category_levels(extra, "Boxe") == conflict_a
    elif case == "all_invalid":
        extra = {"favorite_data": {"category_trail": "TV/Casti"},
                 "data-category-trail": "TV/Telefoane"}
        assert select_product_category_levels(extra, "Boxe") == []
    elif case == "final_beats_temporary":
        decision = decide_category_evidence([
            _final(conflict_a), _temporary(conflict_b)])
        assert decision.status == CategoryEvidenceStatus.FINAL
        assert list(decision.evidence.levels) == conflict_a
    else:
        decision = decide_category_evidence([_final(conflict_a), _final(conflict_b)])
        assert decision.status == CategoryEvidenceStatus.FINAL_CONFLICTED
        assert decision.evidence is None


def _product(number, category_url, trail_key="favorite", levels=None):
    levels = levels or ["Catalog", f"Own-{number}", "Boxe"]
    extra = {}
    if trail_key == "favorite":
        extra["favorite_data"] = {"category_trail": "/".join(levels)}
    elif trail_key == "card":
        extra["data-category-trail"] = "/".join(levels)
    return ProductItem(
        category_name="Boxe", category_url=category_url,
        source_page_url=category_url, page_number=number,
        position_in_page=number, product_id=str(number), pnk=f"PNK{number}",
        offer_id=f"OFF{number}", title=f"Produs {number}",
        product_url=f"https://www.emag.ro/produs-{number}/pd/PNK{number}/",
        price_current=100.25 + number, rating=4.5, review_count=number,
        extra=extra,
    )


def _card(number, levels):
    favorite = html.escape(json.dumps({
        "productid": str(number), "pnk": f"PNK{number}",
        "offerid": f"OFF{number}", "product_name": f"Produs {number}",
        "currency": "RON", "price": 100.25 + number,
        "category_trail": "/".join(levels),
    }), quote=True)
    return f'''<div class="card-item card-standard js-product-data" data-product-id="{number}"
      data-name="Produs {number}" data-position="{number}"
      data-url="https://www.emag.ro/produs-{number}/pd/PNK{number}/">
      <button class="add-to-favorites" data-product="{favorite}"></button>
      <p class="product-new-price">{100 + number},25 Lei</p></div>'''


def _listing(number, public_levels, own_levels, next_page=None):
    next_link = f'<link rel="next" href="/boxe/p{next_page}/c">' if next_page else ""
    return (f"<html><head><title>Boxe</title>{next_link}</head><body>" +
            _linked_breadcrumb(public_levels) + _card(number, own_levels) +
            "</body></html>")


@pytest.fixture(scope="module")
def cross_page_run(tmp_path_factory):
    output = tmp_path_factory.mktemp("v221_round2_cross_page")
    pages = {
        1: _listing(1, FINAL_B, ["Own", "Page-1", "Boxe"], 2),
        2: _listing(2, FINAL_B, ["Own", "Page-2", "Boxe"], 3),
        3: _listing(3, FINAL_A, ["Own", "Page-3", "Boxe"]),
    }

    class CountingCrawler(EmagCrawler):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.soup_calls = 0
            self.evidence_calls = 0

        def _parse_html_once(self, body):
            self.soup_calls += 1
            return super()._parse_html_once(body)

        def _get_or_extract_category_evidence(self, *args, **kwargs):
            self.evidence_calls += 1
            return super()._get_or_extract_category_evidence(*args, **kwargs)

    crawler = CountingCrawler(
        str(output), download_images=False, page_workers=3,
        category_workers=1, max_in_flight=3)
    requests = []

    def fake_fetch(url):
        number = 1
        for candidate in (2, 3):
            if f"/p{candidate}/" in url:
                number = candidate
        requests.append(number)
        # Complete out of order while keeping the observed logical page set 1..3.
        if number == 2:
            time.sleep(0.03)
        return FetchResult(
            html=pages[number], status=200, request_url=url, final_url=url,
            request_call_started=True, response_received=True)

    crawler._fetch_page = fake_fetch
    crawler.crawl_all_categories(
        [{"name": "Boxe", "url": BOXE_URL, "enabled": True}], max_pages=3)
    summary = crawler.finalize()
    json_rows = json.loads((output / "products.json").read_text(encoding="utf-8"))
    with (output / "products.csv").open(encoding="utf-8-sig", newline="") as handle:
        csv_rows = list(csv.DictReader(handle))
    workbook = load_workbook(output / "products.xlsx", data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    xlsx_rows = [dict(zip(headers, [cell.value for cell in row]))
                 for row in sheet.iter_rows(min_row=2)]
    workbook.close()
    return {
        "crawler": crawler, "requests": requests, "summary": summary,
        "json": json_rows, "csv": csv_rows, "xlsx": xlsx_rows,
    }


REGISTRY_CRAWLER_CASES = (
    "temporary_long_final_short", "final_short_temporary_long",
    "temporary_short_final_long", "final_long_temporary_short",
    "consistent_finals_choose_longer", "finals_conflict_state",
    "final_conflict_get_none", "final_conflict_sticky_temporary",
    "final_conflict_sticky_third_final", "temporary_conflict_final_upgrade",
    "second_page_finds_final_conflict", "third_page_finds_final_conflict",
    "observation_budget_three", "no_additional_requests", "one_soup_per_page",
    "one_evidence_extract_per_page", "page_one_product_falls_back",
    "page_two_product_falls_back", "later_product_falls_back",
    "concurrent_completion_is_stable", "random_registration_is_stable",
    "different_urls_isolated", "conflict_does_not_pollute_other",
    "real_crawler_chain_final_conflicted",
)


@pytest.mark.parametrize("case", REGISTRY_CRAWLER_CASES)
def test_d_registry_and_cross_page_crawler(case, cross_page_run, tmp_path):
    """D30-D53: atomic states and logical three-page observation are deterministic."""
    registry = CategoryLevelRegistry()
    if case == "temporary_long_final_short":
        registry.register(BOXE_URL, _temporary(FINAL_B)); registry.register(BOXE_URL, _final(FINAL_A))
        assert list(registry.get(BOXE_URL).levels) == FINAL_A
    elif case == "final_short_temporary_long":
        registry.register(BOXE_URL, _final(FINAL_A)); registry.register(BOXE_URL, _temporary(FINAL_B))
        assert list(registry.get(BOXE_URL).levels) == FINAL_A
    elif case == "temporary_short_final_long":
        registry.register(BOXE_URL, _temporary(FINAL_A)); registry.register(BOXE_URL, _final(FINAL_B))
        assert list(registry.get(BOXE_URL).levels) == FINAL_B
    elif case == "final_long_temporary_short":
        registry.register(BOXE_URL, _final(FINAL_B)); registry.register(BOXE_URL, _temporary(FINAL_A))
        assert list(registry.get(BOXE_URL).levels) == FINAL_B
    elif case == "consistent_finals_choose_longer":
        registry.register(BOXE_URL, _final(["TV", "Boxe"])); registry.register(BOXE_URL, _final(FINAL_A))
        assert list(registry.get(BOXE_URL).levels) == FINAL_A
    elif case in {"finals_conflict_state", "final_conflict_get_none",
                  "final_conflict_sticky_temporary", "final_conflict_sticky_third_final"}:
        registry.register(BOXE_URL, _final(FINAL_A)); registry.register(BOXE_URL, _final(FINAL_B))
        if case == "final_conflict_sticky_temporary":
            registry.register(BOXE_URL, _temporary(["Audio", "Boxe"]))
        if case == "final_conflict_sticky_third_final":
            registry.register(BOXE_URL, _final(["Other", "Audio", "Boxe"]))
        assert registry.get_state(BOXE_URL) == CategoryEvidenceStatus.FINAL_CONFLICTED
        assert registry.get(BOXE_URL) is None
    elif case == "temporary_conflict_final_upgrade":
        registry.register(BOXE_URL, _temporary(FINAL_A)); registry.register(BOXE_URL, _temporary(FINAL_B))
        assert registry.get_state(BOXE_URL) == CategoryEvidenceStatus.TEMPORARY_CONFLICTED
        registry.register(BOXE_URL, _final(FINAL_A))
        assert registry.get_state(BOXE_URL) == CategoryEvidenceStatus.FINAL
    elif case in {"second_page_finds_final_conflict", "third_page_finds_final_conflict",
                  "observation_budget_three"}:
        crawler = EmagCrawler(str(tmp_path / case), download_images=False)
        pages = [FINAL_A, FINAL_B] if case == "second_page_finds_final_conflict" else [FINAL_A, FINAL_A, FINAL_B]
        for page_number, levels in enumerate(pages, 1):
            crawler._get_or_extract_category_evidence(
                _soup(_linked_breadcrumb(levels)), "Boxe", BOXE_URL,
                page_number=page_number)
        if case == "observation_budget_three":
            crawler._get_or_extract_category_evidence(
                _soup(_linked_breadcrumb(["Ignored", "Boxe"])), "Boxe", BOXE_URL,
                page_number=4)
            assert crawler._category_level_upgrade_checks["emag.ro/boxe/c"] == 3
        assert crawler.exporters.get_category_level_state(BOXE_URL) == (
            CategoryEvidenceStatus.FINAL_CONFLICTED)
    elif case == "no_additional_requests":
        assert sorted(cross_page_run["requests"]) == [1, 2, 3]
    elif case == "one_soup_per_page":
        assert cross_page_run["crawler"].soup_calls == 3
    elif case == "one_evidence_extract_per_page":
        assert cross_page_run["crawler"].evidence_calls == 3
    elif case in {"page_one_product_falls_back", "page_two_product_falls_back",
                  "later_product_falls_back"}:
        index = {"page_one_product_falls_back": 0, "page_two_product_falls_back": 1,
                 "later_product_falls_back": 2}[case]
        row = cross_page_run["json"][index]
        assert [row["一级类"], row["二级类"], row["三级类"]] == [
            "Own", f"Page-{index + 1}", "Boxe"]
    elif case in {"concurrent_completion_is_stable", "random_registration_is_stable"}:
        for seed in range(12):
            local = CategoryLevelRegistry()
            candidates = [_final(FINAL_A), _final(FINAL_B)] * 8
            random.Random(seed).shuffle(candidates)
            with ThreadPoolExecutor(max_workers=8) as executor:
                list(executor.map(lambda item: local.register(BOXE_URL, item), candidates))
            assert local.get_state(BOXE_URL) == CategoryEvidenceStatus.FINAL_CONFLICTED
            assert local.get(BOXE_URL) is None
    elif case == "different_urls_isolated":
        registry.register(BOXE_URL, _final(FINAL_A))
        registry.register("https://www.emag.ro/casti/c", _final(["Audio", "Casti"]))
        assert registry.get(BOXE_URL).levels[-1] == "Boxe"
        assert registry.get("https://www.emag.ro/casti/c").levels[-1] == "Casti"
    elif case == "conflict_does_not_pollute_other":
        registry.register(BOXE_URL, _final(FINAL_A)); registry.register(BOXE_URL, _final(FINAL_B))
        registry.register("https://www.emag.ro/casti/c", _final(["Audio", "Casti"]))
        assert registry.get(BOXE_URL) is None
        assert list(registry.get("https://www.emag.ro/casti/c").levels) == ["Audio", "Casti"]
    else:
        assert cross_page_run["crawler"].exporters.get_category_level_state(BOXE_URL) == (
            CategoryEvidenceStatus.FINAL_CONFLICTED)
        assert len(cross_page_run["json"]) == 3


class _SessionHandler(http.server.BaseHTTPRequestHandler):
    lock = threading.Lock()
    count = 0
    cookies = []

    def do_GET(self):
        with type(self).lock:
            type(self).count += 1
            type(self).cookies.append(self.headers.get("Cookie", ""))
        body = b"<html><body>ok</body></html>"
        self.send_response(500 if self.path.startswith("/error") else 200)
        if not self.headers.get("Cookie"):
            self.send_header("Set-Cookie", "round2_session=kept; Path=/")
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, _format, *_args):
        return


class _ThreadingServer(socketserver.ThreadingMixIn, socketserver.TCPServer):
    allow_reuse_address = True
    daemon_threads = True


@pytest.fixture(scope="module")
def session_server():
    proxy_names = ("ALL_PROXY", "HTTP_PROXY", "HTTPS_PROXY", "all_proxy",
                   "http_proxy", "https_proxy")
    saved = {name: os.environ.get(name) for name in proxy_names}
    for name in proxy_names:
        os.environ.pop(name, None)
    server = _ThreadingServer(("127.0.0.1", 0), _SessionHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_address[1]}"
    finally:
        server.shutdown(); server.server_close(); thread.join(timeout=2)
        for name, value in saved.items():
            if value is not None:
                os.environ[name] = value


SESSION_CASES = (
    "same_generation_same_thread", "different_threads", "cookie_continuity",
    "finalize_closes_generation", "finalize_is_idempotent", "new_object_after_finalize",
    "new_client_active", "new_client_requests", "same_worker_thread_local_invalidated",
    "two_request_finalize_rounds", "no_session_leak_or_double_close",
)


@pytest.mark.parametrize("case", SESSION_CASES)
def test_e_session_generation_lifecycle(case, session_server, tmp_path):
    """E54-E64: generation invalidates closed thread-local clients safely."""
    crawler = EmagCrawler(str(tmp_path / case), download_images=False)
    crawler._session_config["retries"] = 1
    url = session_server + "/ok/c"
    if case == "same_generation_same_thread":
        assert crawler._get_client() is crawler._get_client()
    elif case == "different_threads":
        barrier = threading.Barrier(2)
        def acquire(_):
            client = crawler._get_client(); barrier.wait(); return id(client)
        with ThreadPoolExecutor(max_workers=2) as executor:
            assert len(set(executor.map(acquire, range(2)))) == 2
    elif case == "cookie_continuity":
        start = len(_SessionHandler.cookies)
        client = crawler._get_client(); client.get(url); client.get(url)
        assert _SessionHandler.cookies[start] == ""
        assert "round2_session=kept" in _SessionHandler.cookies[start + 1]
    elif case == "finalize_closes_generation":
        crawler._get_client(); old = crawler._session_generation; crawler.finalize()
        assert crawler._session_generation == old + 1 and crawler._all_sessions == []
    elif case == "finalize_is_idempotent":
        crawler._get_client(); crawler.finalize(); crawler.finalize()
        assert crawler._all_sessions == []
    elif case == "new_object_after_finalize":
        first = crawler._get_client(); crawler.finalize(); second = crawler._get_client()
        assert first is not second
    elif case == "new_client_active":
        crawler._get_client(); crawler.finalize(); client = crawler._get_client()
        assert crawler._session_is_active(crawler._thread_local.mgr, client)
    elif case == "new_client_requests":
        crawler._get_client(); crawler.finalize(); assert crawler._get_client().get(url).status == 200
    elif case == "same_worker_thread_local_invalidated":
        with ThreadPoolExecutor(max_workers=1) as executor:
            first = executor.submit(lambda: id(crawler._get_client())).result()
            crawler.finalize()
            second = executor.submit(lambda: id(crawler._get_client())).result()
        assert first != second
    elif case == "two_request_finalize_rounds":
        identities = []
        for _ in range(2):
            client = crawler._get_client(); identities.append(id(client))
            assert client.get(url).status == 200; crawler.finalize()
        assert identities[0] != identities[1]
    else:
        managers = []
        for _ in range(2):
            crawler._get_client(); managers.extend(m for m, _c in crawler._all_sessions)
            crawler.finalize(); crawler.finalize()
        assert crawler._all_sessions == []
        assert all(not manager._is_alive for manager in managers)
    crawler._close_all_sessions()


DIAGNOSTIC_CASES = (
    "creation_failure", "before_get_failure", "inside_get_failure",
    "response_received", "retries_zero", "retries_negative",
    "retries_one_single_call", "secrets_redacted",
)


class _FakePage:
    html_content = "<html><body>ok</body></html>"
    status = 200
    url = BOXE_URL
    headers = {"Content-Type": "text/html"}
    history = []


@pytest.mark.parametrize("case", DIAGNOSTIC_CASES)
def test_f_request_phase_diagnostics(case, session_server, tmp_path):
    """F65-F72: request_call_started and response_received reflect real phases."""
    crawler = EmagCrawler(str(tmp_path / case), download_images=False)
    if case == "creation_failure":
        crawler._get_client = lambda: (_ for _ in ()).throw(RuntimeError("create failed"))
        result = crawler._fetch_page(BOXE_URL)
        assert not result.request_call_started and not result.response_received
        assert "phase=session_creation" in result.error_detail
    elif case == "before_get_failure":
        client = object(); crawler._get_client = lambda: client
        crawler._validate_client_before_request = lambda _client: (_ for _ in ()).throw(RuntimeError("inactive"))
        result = crawler._fetch_page(BOXE_URL)
        assert not result.request_call_started and "phase=before_client_get" in result.error_detail
    elif case == "inside_get_failure":
        class Broken:
            def get(self, _url): raise OSError("network down")
        crawler._get_client = lambda: Broken(); crawler._validate_client_before_request = lambda _client: None
        result = crawler._fetch_page(BOXE_URL)
        assert result.request_call_started and not result.response_received
        assert "phase=inside_client_get" in result.error_detail
    elif case == "response_received":
        class Working:
            def get(self, _url): return _FakePage()
        crawler._get_client = lambda: Working(); crawler._validate_client_before_request = lambda _client: None
        result = crawler._fetch_page(BOXE_URL)
        assert result.request_call_started and result.response_received and result.status == 200
    elif case in {"retries_zero", "retries_negative"}:
        crawler._session_config["retries"] = 0 if case == "retries_zero" else -1
        before = _SessionHandler.count
        result = crawler._fetch_page(session_server + "/ok/c")
        assert result.error_type == "ValueError" and not result.request_call_started
        assert _SessionHandler.count == before
    elif case == "retries_one_single_call":
        crawler._session_config["retries"] = 1; before = _SessionHandler.count
        result = crawler._fetch_page(session_server + "/error/c")
        assert result.status == 500 and result.response_received
        assert _SessionHandler.count == before + 1
    else:
        secret = "Cookie: private; Token=secretvalue; Authorization: Bearer hidden"
        crawler._get_client = lambda: (_ for _ in ()).throw(RuntimeError(secret))
        result = crawler._fetch_page(BOXE_URL)
        assert "private" not in result.error_detail
        assert "secretvalue" not in result.error_detail
        assert "Bearer hidden" not in result.error_detail
        assert result.error_detail.count("[REDACTED]") == 3
    crawler._close_all_sessions()


@pytest.fixture(scope="module")
def format_round_trip(tmp_path_factory):
    output = tmp_path_factory.mktemp("v221_round2_formats")
    exporter = Exporters(str(output))
    urls = {
        "same": "https://www.emag.ro/same-page/c",
        "cross": "https://www.emag.ro/cross-page/c",
        "card": "https://www.emag.ro/card-fallback/c",
        "invalid": "https://www.emag.ro/no-path/c",
        "five": "https://www.emag.ro/five/c",
    }
    exporter._category_levels.register_conflict(urls["same"], final=False)
    exporter.register_category_levels(urls["cross"], _final(["Wrong-A", "Boxe"]))
    exporter.register_category_levels(urls["cross"], _final(["Wrong-B", "Boxe"]))
    exporter.register_category_levels(
        urls["five"], _final(["L1", "L2", "L3", "L4", "Boxe"]))
    products = [
        _product(1, urls["same"], "favorite", ["Own", "Same", "Boxe"]),
        _product(2, urls["cross"], "favorite", ["Own", "Early-1", "Boxe"]),
        _product(3, urls["cross"], "favorite", ["Own", "Early-2", "Boxe"]),
        _product(4, urls["cross"], "favorite", ["Own", "Late", "Boxe"]),
        _product(5, urls["card"], "card", ["TV", "Audio", "Hi-Fi", "Boxe"]),
        _product(6, urls["invalid"], "none"),
        _product(7, urls["five"], "favorite", ["Fallback", "Boxe"]),
    ]
    before = [copy.deepcopy(product.to_dict()) for product in products]
    baseline = [product_to_output_dict(item) for item in before]
    exporter.add_products(products); exporter.finalize()
    json_path = output / "products.json"; csv_path = output / "products.csv"
    xlsx_path = output / "products.xlsx"
    json_rows = json.loads(json_path.read_text(encoding="utf-8"))
    raw_csv = csv_path.read_bytes()
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        csv_rows = list(csv.DictReader(handle))
    workbook = load_workbook(xlsx_path, data_only=True); sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    xlsx_rows = [dict(zip(headers, [cell.value for cell in row]))
                 for row in sheet.iter_rows(min_row=2)]
    workbook.close()
    return {
        "products": products, "before": before, "baseline": baseline,
        "json": json_rows, "csv": csv_rows, "xlsx": xlsx_rows,
        "headers": headers, "raw_csv": raw_csv,
        "hashes": {path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                   for path in (json_path, csv_path, xlsx_path)},
    }


FORMAT_CASES = (
    "same_page_temporary_conflict", "cross_page_final_conflict",
    "third_page_early_product", "favorite_fallback", "card_fallback",
    "all_invalid_no_guess", "dynamic_three_four_five", "json_omits_empty",
    "non_category_values", "extra_unchanged", "json_extra_object", "csv_bom",
    "xlsx_numeric_types", "count_order_dedup",
)


@pytest.mark.parametrize("case", FORMAT_CASES)
def test_g_production_three_format_round_trip(case, format_round_trip):
    """G73-G86: real Exporters.finalize output honors conflict fallback boundaries."""
    data = format_round_trip; rows = data["json"]
    if case == "same_page_temporary_conflict":
        assert [rows[0][f] for f in CATEGORY_LEVEL_FIELDS[:3]] == ["Own", "Same", "Boxe"]
    elif case == "cross_page_final_conflict":
        assert [rows[1][f] for f in CATEGORY_LEVEL_FIELDS[:3]] == ["Own", "Early-1", "Boxe"]
    elif case == "third_page_early_product":
        assert [rows[2][f] for f in CATEGORY_LEVEL_FIELDS[:3]] == ["Own", "Early-2", "Boxe"]
    elif case == "favorite_fallback":
        assert rows[3]["二级类"] == "Late"
    elif case == "card_fallback":
        assert [rows[4][f] for f in CATEGORY_LEVEL_FIELDS[:4]] == ["TV", "Audio", "Hi-Fi", "Boxe"]
    elif case == "all_invalid_no_guess":
        assert not any(field in rows[5] for field in CATEGORY_LEVEL_FIELDS)
    elif case == "dynamic_three_four_five":
        assert data["headers"][3:8] == list(CATEGORY_LEVEL_FIELDS)
        assert [rows[6][f] for f in CATEGORY_LEVEL_FIELDS] == ["L1", "L2", "L3", "L4", "Boxe"]
    elif case == "json_omits_empty":
        assert "四级类" not in rows[0] and "五级类" not in rows[0]
    elif case == "non_category_values":
        for output, baseline in zip(rows, data["baseline"]):
            assert {k: v for k, v in output.items() if k not in CATEGORY_LEVEL_FIELDS} == {
                k: v for k, v in baseline.items() if k not in CATEGORY_LEVEL_FIELDS}
    elif case == "extra_unchanged":
        assert [product.to_dict() for product in data["products"]] == data["before"]
    elif case == "json_extra_object":
        assert all(isinstance(row.get("扩展信息", {}), dict) for row in rows)
    elif case == "csv_bom":
        assert data["raw_csv"].startswith(b"\xef\xbb\xbf")
    elif case == "xlsx_numeric_types":
        assert all(isinstance(row["前端价格"], float) for row in data["xlsx"])
        assert all(isinstance(row["评论分数"], float) for row in data["xlsx"])
        assert all(isinstance(row["评价数量"], int) for row in data["xlsx"])
    else:
        assert len(rows) == len(data["csv"]) == len(data["xlsx"]) == 7
        assert [row["PNK码"] for row in rows] == [f"PNK{i}" for i in range(1, 8)]
        assert len({row["PNK码"] for row in rows}) == 7
